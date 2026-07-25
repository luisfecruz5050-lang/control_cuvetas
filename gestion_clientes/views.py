from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponseRedirect, HttpResponse
from django.urls import reverse_lazy
from django.views import View
from django.views.generic import ListView, CreateView, UpdateView, DeleteView, DetailView
from braces.views import PermissionRequiredMixin, MultiplePermissionsRequiredMixin
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.messages.views import SuccessMessageMixin
from django.contrib import messages
from .models import Clientes
from .forms import ClientesForm
from django.db.models.deletion import ProtectedError
from datetime import datetime
from django.utils import timezone
from django.core.paginator import Paginator, PageNotAnInteger, EmptyPage
import openpyxl
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from gestion_cuvetas.models import ControlCuvetas



class DescargarClienteExcel(LoginRequiredMixin, PermissionRequiredMixin, View):
    """Genera y descarga el informe Excel de un cliente con su estado de cuvetas."""
    permission_required = 'gestion_clientes.excel_clientes'

    def get(self, request, pk, *args, **kwargs):
        cliente = get_object_or_404(Clientes, pk=pk)
        controles = ControlCuvetas.objects.filter(cliente=cliente).order_by('-fecha_registro')

        # ── Cálculos de cuvetas por estado ────────────────────
        def sumar(qs, campo):
            return sum([getattr(c, campo) or 0 for c in qs])

        entregado_qs   = controles.filter(estado='EN')
        parcial_qs     = controles.filter(estado='PA')
        devuelto_qs    = controles.filter(estado='DE')
        pendiente_qs   = controles.filter(estado='PE')

        total_grandes   = sumar(controles, 'cuvetas_grandes')
        total_pequenas  = sumar(controles, 'cuvetas_pequenas')
        total_cuvetas   = total_grandes + total_pequenas

        entregado_total  = sumar(entregado_qs, 'cuvetas_grandes') + sumar(entregado_qs, 'cuvetas_pequenas')
        parcial_total    = sumar(parcial_qs, 'cuvetas_grandes') + sumar(parcial_qs, 'cuvetas_pequenas')
        devuelto_total   = sumar(devuelto_qs, 'cuvetas_grandes') + sumar(devuelto_qs, 'cuvetas_pequenas')
        pendiente_total  = sumar(pendiente_qs, 'cuvetas_grandes') + sumar(pendiente_qs, 'cuvetas_pequenas')

        nombre_cliente = f"{cliente.nombre or ''} {cliente.apellido or ''}".strip() or (cliente.nombre_razon or 'Sin nombre')
        identificacion = cliente.identificacion or '—'
        telefono       = cliente.telefono or '—'
        direccion      = cliente.direccion or '—'
        ciudad         = cliente.ciudad or '—'
        negocio        = cliente.negocio or '—'

        # ══════════════════════════════════════════════
        # PALETA (misma que tu ejemplo)
        # ══════════════════════════════════════════════
        COLOR_HEADER = "1F4E79"
        COLOR_COL    = "BDD7EE"
        COLOR_PAR    = "DEEAF1"
        COLOR_TOT    = "2E75B6"
        BLANCO       = "FFFFFF"
        GRIS_LABEL   = "EFF6FF"

        COLOR_OK     = "2E7D32"   # verde - devuelto completo
        COLOR_WARN   = "F9A825"   # amarillo - parcial
        COLOR_INFO   = "1565C0"   # azul - entregado
        COLOR_ALERT  = "C62828"   # rojo - pendiente

        # ══════════════════════════════════════════════
        # HELPERS
        # ══════════════════════════════════════════════
        def fl(hex_color):
            return PatternFill("solid", fgColor=hex_color)

        fl_blanco      = fl(BLANCO)
        fl_header      = fl(COLOR_HEADER)
        fl_col         = fl(COLOR_COL)
        fl_gris        = fl(GRIS_LABEL)
        fl_tot         = fl(COLOR_TOT)

        bd_std = Border(
            left   = Side(style="thin", color="B0BEC5"),
            right  = Side(style="thin", color="B0BEC5"),
            top    = Side(style="thin", color="B0BEC5"),
            bottom = Side(style="thin", color="B0BEC5"),
        )
        bd_none = Border()

        AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
        AL_L = Alignment(horizontal="left",   vertical="center", wrap_text=True)
        AL_R = Alignment(horizontal="right",  vertical="center")

        f_titulo    = Font(name="Arial", bold=True,  size=14, color=BLANCO)
        f_gen       = Font(name="Arial", italic=True,size=9,  color="C7D2FE")
        f_subtitulo = Font(name="Arial", bold=True,  size=10, color=COLOR_HEADER)
        f_col_hdr   = Font(name="Arial", bold=True,  size=10, color=COLOR_HEADER)
        f_lbl       = Font(name="Arial", bold=True,  size=10, color="1E3A5F")
        f_val       = Font(name="Arial",             size=10, color="0F172A")
        f_dato      = Font(name="Arial",             size=10, color="1E293B")
        f_dato_bold = Font(name="Arial", bold=True,  size=10, color="0F172A")
        f_tot_sm    = Font(name="Arial", bold=True,  size=10, color=BLANCO)
        f_tot_lg    = Font(name="Arial", bold=True,  size=11, color=BLANCO)
        f_footer    = Font(name="Arial", italic=True,size=8,  color="94A3B8")

        def pintar_rango(ws, fila, col_ini, col_fin, fill_obj,
                         font=None, border=None, alineacion=None, valor=None):
            for col in range(col_ini, col_fin + 1):
                c = ws.cell(row=fila, column=col)
                c.fill = fill_obj
                if border:
                    c.border = border
            if col_ini != col_fin:
                ws.merge_cells(start_row=fila, start_column=col_ini,
                               end_row=fila,   end_column=col_fin)
            c_main = ws.cell(row=fila, column=col_ini)
            if valor      is not None: c_main.value      = valor
            if font       is not None: c_main.font       = font
            if alineacion is not None: c_main.alignment  = alineacion
            return c_main

        def fila_blanca(ws, fila, alto=8, columnas=7):
            ws.row_dimensions[fila].height = alto
            for col in range(1, columnas + 1):
                c = ws.cell(row=fila, column=col)
                c.fill   = fl_blanco
                c.border = bd_none

        ESTADO_MAP = {
            'EN': ('Entregado', COLOR_INFO),
            'PA': ('Parcialmente devuelto', COLOR_WARN),
            'DE': ('Devuelto completo', COLOR_OK),
            'PE': ('Pendiente', COLOR_ALERT),
        }

        # ══════════════════════════════════════════════
        # WORKBOOK
        # ══════════════════════════════════════════════
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Informe Cliente"

        # 7 columnas: #, Fecha, Grandes, Pequeñas, Total, Estado, Observaciones
        anchos = {"A": 6, "B": 14, "C": 12, "D": 14, "E": 10, "F": 20, "G": 30}
        for col_letra, ancho in anchos.items():
            ws.column_dimensions[col_letra].width = ancho

        fila = 1

        # ─────────────────────────────────────────────
        # 1. ENCABEZADO
        # ─────────────────────────────────────────────
        ws.row_dimensions[fila].height = 30
        pintar_rango(ws, fila, 1, 7, fl_header, f_titulo, bd_none, AL_C,
                     "INFORME DE CLIENTE - CONTROL DE CUVETAS")
        fila += 1

        ws.row_dimensions[fila].height = 24
        pintar_rango(ws, fila, 1, 7, fl_header, f_titulo, bd_none, AL_C,
                     nombre_cliente.upper())
        fila += 1

        ws.row_dimensions[fila].height = 16
        pintar_rango(ws, fila, 1, 7, fl_header, f_gen, bd_none, AL_C,
                     f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}")
        fila += 1

        fila_blanca(ws, fila, 10)
        fila += 1

        # ─────────────────────────────────────────────
        # 2. DATOS DEL CLIENTE
        # ─────────────────────────────────────────────
        def escribir_info(ws, fila, label, valor):
            ws.row_dimensions[fila].height = 18
            pintar_rango(ws, fila, 1, 2, fl_gris,   f_lbl, bd_std, AL_L, label)
            pintar_rango(ws, fila, 3, 7, fl_blanco, f_val, bd_std, AL_L, valor)
            return fila + 1

        info_fields = [
            ("Identificación:", identificacion),
            ("Teléfono:",       telefono),
            ("Negocio:",        negocio),
            ("Dirección:",      direccion),
            ("Ciudad:",         ciudad),
        ]
        for label, valor in info_fields:
            fila = escribir_info(ws, fila, label, valor)

        fila_blanca(ws, fila, 10)
        fila += 1

        # ─────────────────────────────────────────────
        # 3. RESUMEN GENERAL DE CUVETAS
        # ─────────────────────────────────────────────
        ws.row_dimensions[fila].height = 22
        pintar_rango(ws, fila, 1, 7, fl_col, f_subtitulo, bd_std, AL_C,
                     "RESUMEN GENERAL DE CUVETAS")
        fila += 1

        resumen = [
            ("Total cuvetas entregadas (histórico)", total_cuvetas,     COLOR_HEADER),
            ("Entregadas (estado Entregado)",         entregado_total,  COLOR_INFO),
            ("Parcialmente devueltas",                parcial_total,    COLOR_WARN),
            ("Devueltas completamente",                devuelto_total,   COLOR_OK),
            ("Pendientes",                              pendiente_total,  COLOR_ALERT),
        ]

        for label, valor, color in resumen:
            ws.row_dimensions[fila].height = 22
            fill_bg = fl(color)
            pintar_rango(ws, fila, 1, 5, fill_bg, f_tot_sm, bd_std, AL_R, label)
            pintar_rango(ws, fila, 6, 7, fill_bg, f_tot_sm, bd_std, AL_C, valor)
            fila += 1

        fila_blanca(ws, fila, 10)
        fila += 1

        # ─────────────────────────────────────────────
        # 4. TÍTULO DETALLE DE MOVIMIENTOS
        # ─────────────────────────────────────────────
        ws.row_dimensions[fila].height = 22
        pintar_rango(ws, fila, 1, 7, fl_col, f_subtitulo, bd_std, AL_C,
                     "DETALLE DE MOVIMIENTOS DE CUVETAS")
        fila += 1

        # ─────────────────────────────────────────────
        # 5. CABECERAS DE TABLA
        # ─────────────────────────────────────────────
        ws.row_dimensions[fila].height = 20
        pintar_rango(ws, fila, 1, 1, fl_col, f_col_hdr, bd_std, AL_C, "#")
        pintar_rango(ws, fila, 2, 2, fl_col, f_col_hdr, bd_std, AL_C, "Fecha")
        pintar_rango(ws, fila, 3, 3, fl_col, f_col_hdr, bd_std, AL_C, "Grandes")
        pintar_rango(ws, fila, 4, 4, fl_col, f_col_hdr, bd_std, AL_C, "Pequeñas")
        pintar_rango(ws, fila, 5, 5, fl_col, f_col_hdr, bd_std, AL_C, "Total")
        pintar_rango(ws, fila, 6, 6, fl_col, f_col_hdr, bd_std, AL_C, "Estado")
        pintar_rango(ws, fila, 7, 7, fl_col, f_col_hdr, bd_std, AL_C, "Observaciones")
        fila += 1

        # ─────────────────────────────────────────────
        # 6. FILAS DE DETALLE
        # ─────────────────────────────────────────────
        if not controles.exists():
            ws.row_dimensions[fila].height = 28
            pintar_rango(ws, fila, 1, 7, fl_blanco, f_dato, bd_std, AL_C,
                         "Este cliente no tiene movimientos de cuvetas registrados.")
            fila += 1
        else:
            for idx, control in enumerate(controles, start=1):
                ws.row_dimensions[fila].height = 22
                bg = fl(COLOR_PAR) if idx % 2 == 0 else fl_blanco

                fecha_str    = control.fecha_registro.strftime('%d/%m/%Y') if control.fecha_registro else '—'
                estado_label, estado_color = ESTADO_MAP.get(control.estado, (control.estado, COLOR_HEADER))

                pintar_rango(ws, fila, 1, 1, bg, f_dato,      bd_std, AL_C, idx)
                pintar_rango(ws, fila, 2, 2, bg, f_dato,      bd_std, AL_C, fecha_str)
                pintar_rango(ws, fila, 3, 3, bg, f_dato,      bd_std, AL_C, control.cuvetas_grandes or 0)
                pintar_rango(ws, fila, 4, 4, bg, f_dato,      bd_std, AL_C, control.cuvetas_pequenas or 0)
                pintar_rango(ws, fila, 5, 5, bg, f_dato_bold, bd_std, AL_C, control.total_cuvetas)

                c_estado = pintar_rango(ws, fila, 6, 6, fl(estado_color), f_tot_sm, bd_std, AL_C, estado_label)

                pintar_rango(ws, fila, 7, 7, bg, f_dato, bd_std, AL_L, control.observaciones or '—')

                fila += 1

        fila_blanca(ws, fila, 6)
        fila += 1

        # ─────────────────────────────────────────────
        # 7. FOOTER
        # ─────────────────────────────────────────────
        ws.row_dimensions[fila].height = 16
        bd_top = Border(top=Side(style="thin", color="D1D5DB"))
        pintar_rango(ws, fila, 1, 7, fl_blanco, f_footer, bd_top, AL_C,
                     "Documento generado automáticamente — no requiere firma física.")

        # ─────────────────────────────────────────────
        # RESPUESTA HTTP
        # ─────────────────────────────────────────────
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="informe_cliente_{cliente.pk}_{timezone.now().strftime("%Y-%m-%d")}.xlsx"'
        )
        wb.save(response)
        return response


####--------Vistas del curso--------####
class ListaClientes(LoginRequiredMixin, MultiplePermissionsRequiredMixin, ListView):
    model = Clientes
    template_name = 'gestion_clientes/listar_clientes.html'
    context_object_name = 'clientes'
    paginate_by = 10
    ordering = ['-fecha_registro']
    permissions = {
        "any": (
            'gestion_clientes.listar_clientes',
            'gestion_clientes.registrar_cliente',
            'gestion_clientes.actualizar_cliente',
            'gestion_clientes.eliminar_cliente',
            'gestion_clientes.detalle_cliente',
            'gestion_clientes.cargar_clientes',    # ← nuevo
        )
    }

    def get_queryset(self):
        queryset = super().get_queryset()
        filtro_nombre         = self.request.GET.get('nombre', '').strip()
        filtro_apellido       = self.request.GET.get('apellido', '').strip()
        filtro_identificacion = self.request.GET.get('identificacion', '').strip()
        filtro_telefono       = self.request.GET.get('telefono', '').strip()
        filtro_email          = self.request.GET.get('email', '').strip()
        filtro_genero         = self.request.GET.get('genero', '').strip()

        if filtro_nombre:
            queryset = queryset.filter(nombre__istartswith=filtro_nombre)
        if filtro_apellido:
            queryset = queryset.filter(apellido__istartswith=filtro_apellido)
        if filtro_identificacion:
            if filtro_identificacion.isdigit():
                queryset = queryset.filter(identificacion__istartswith=filtro_identificacion)
            else:
                queryset = queryset.none()
        if filtro_telefono:
            if filtro_telefono.isdigit():
                queryset = queryset.filter(telefono__istartswith=filtro_telefono)
            else:
                queryset = queryset.none()
        if filtro_email:
            queryset = queryset.filter(email__istartswith=filtro_email)
        if filtro_genero:
            queryset = queryset.filter(genero__istartswith=filtro_genero)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['filtro_nombre']         = self.request.GET.get('nombre', '')
        context['filtro_apellido']       = self.request.GET.get('apellido', '')
        context['filtro_identificacion'] = self.request.GET.get('identificacion', '')
        context['filtro_telefono']       = self.request.GET.get('telefono', '')
        context['filtro_email']          = self.request.GET.get('email', '')
        context['filtro_genero']         = self.request.GET.get('genero', '')
        querystring = self.request.GET.copy()
        if 'page' in querystring:
            querystring.pop('page')
        context['querystring'] = '&' + querystring.urlencode() if querystring else ''
        return context


# ─────────────────────────────────────────────────────────────────────────────
# CARGUE MASIVO DESDE EXCEL
# ─────────────────────────────────────────────────────────────────────────────

# Cabeceras esperadas en el Excel (en minúsculas, sin espacios)
_CABECERAS_EXCEL = [
    'nombre', 'apellido', 'nombre_razon', 'tipo_identificacion',
    'identificacion', 'telefono', 'email', 'genero',
    'direccion', 'barrio', 'negocio', 'ciudad', 'departamento',
]


class CargueClientesExcel(LoginRequiredMixin, PermissionRequiredMixin, View):
    """
    Recibe un archivo Excel (.xlsx) y crea registros de Clientes masivamente.
    El POST se envía desde el modal en listar_clientes.html.
    """
    permission_required = 'gestion_clientes.cargar_clientes'

    def get(self, request, *args, **kwargs):
        # El formulario vive en el modal de listar_clientes; no hay página propia
        return redirect('listar_clientes')

    def post(self, request, *args, **kwargs):
        archivo = request.FILES.get('archivo_excel')

        # ── Validaciones de archivo ──────────────────────────────────────────
        if not archivo:
            messages.error(request, "No se seleccionó ningún archivo.")
            return redirect('listar_clientes')

        if not archivo.name.lower().endswith(('.xlsx', '.xls')):
            messages.error(request, "El archivo debe ser un Excel válido (.xlsx o .xls).")
            return redirect('listar_clientes')

        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            ws = wb.active
        except Exception as e:
            messages.error(request, f"No se pudo leer el archivo: {str(e)}")
            return redirect('listar_clientes')

        filas = list(ws.iter_rows(values_only=True))

        if not filas:
            messages.error(request, "El archivo está vacío.")
            return redirect('listar_clientes')

        # ── Validar cabeceras ────────────────────────────────────────────────
        cabeceras_archivo = [
            str(c).strip().lower() if c is not None else ''
            for c in filas[0]
        ]
        cabeceras_faltantes = [h for h in _CABECERAS_EXCEL if h not in cabeceras_archivo]
        if cabeceras_faltantes:
            messages.error(
                request,
                f"El archivo no tiene las columnas requeridas: {', '.join(cabeceras_faltantes)}. "
                "Descarga la plantilla para ver el formato correcto."
            )
            return redirect('listar_clientes')

        # Mapeo nombre_columna → índice
        idx = {h: cabeceras_archivo.index(h) for h in _CABECERAS_EXCEL}

        # ── Procesar filas ───────────────────────────────────────────────────
        creados  = 0
        errores  = []           # lista de strings con descripción del error
        omitidas = 0            # filas completamente vacías

        for num_fila, fila in enumerate(filas[1:], start=2):

            def col(nombre):
                """Devuelve el valor de la celda como string limpio o ''."""
                val = fila[idx[nombre]]
                return str(val).strip() if val is not None else ''

            # Omitir filas vacías
            if all(c is None or str(c).strip() == '' for c in fila):
                omitidas += 1
                continue

            # ── Obtener datos del Excel ──────────────────────────────────────
            nombre = col('nombre') or None
            apellido = col('apellido') or None
            tipo_id = col('tipo_identificacion') or None
            identificacion = col('identificacion') or None
            telefono = col('telefono') or None
            email = col('email') or None
            genero = col('genero') or None

            nombre_razon = col('nombre_razon') or None
            direccion = col('direccion') or None
            barrio = col('barrio') or None
            negocio = col('negocio') or None
            ciudad = col('ciudad') or None
            departamento = col('departamento') or None

            # ── Crear registro ───────────────────────────────────────────────
            try:
                Clientes.objects.create(
                    nombre=nombre,
                    apellido=apellido,
                    nombre_razon=nombre_razon,
                    tipo_identificacion=tipo_id,
                    identificacion=identificacion,
                    telefono=telefono,
                    email=email,
                    genero=genero,
                    direccion=direccion,
                    barrio=barrio,
                    negocio=negocio,
                    ciudad=ciudad,
                    departamento=departamento,
                )
                creados += 1
            except Exception as e:
                errores.append(f"Fila {num_fila}: error al guardar — {str(e)}")

        # ── Mensajes de resultado ────────────────────────────────────────────
        if creados:
            messages.success(request, f"{creados} cliente(s) cargado(s) exitosamente.")

        if errores:
            # Mostramos máximo 10 errores para no saturar la UI
            for err in errores[:10]:
                messages.warning(request, err)
            if len(errores) > 10:
                messages.warning(
                    request,
                    f"… y {len(errores) - 10} error(es) adicionales. Corrige el archivo y vuelve a cargarlo."
                )

        if not creados and not errores:
            messages.info(request, "No se encontraron filas con datos para cargar.")

        return redirect('listar_clientes')


class DescargarPlantillaClientes(LoginRequiredMixin, PermissionRequiredMixin, View):
    """
    Genera y descarga un Excel vacío con las columnas requeridas
    y una fila de ejemplo para orientar al usuario.
    """
    permission_required = 'gestion_clientes.cargar_clientes'

    def get(self, request, *args, **kwargs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clientes"

        # ── Cabeceras ────────────────────────────────────────────────────────
        ws.append(_CABECERAS_EXCEL)

        # Estilo de cabecera (fondo oscuro, texto blanco, negrita)
        from openpyxl.styles import PatternFill, Font, Alignment
        header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = Alignment(horizontal="center")

        # ── Fila de ejemplo ──────────────────────────────────────────────────
        ws.append([
            'Juan',           # nombre
            'Pérez',          # apellido
            'Tienda JuanP',   # nombre_razon  (opcional)
            'CC',             # tipo_identificacion
            '1234567890',     # identificacion
            '3001234567',     # telefono
            'juan@ejemplo.com',# email
            'M',              # genero
            'Calle 5 # 10-20',# direccion     (opcional)
            'El Centro',      # barrio        (opcional)
            'Droguería',      # negocio       (opcional)
            'Cali',           # ciudad        (opcional)
            'Valle del Cauca',# departamento  (opcional)
        ])

        # ── Ancho automático de columnas ─────────────────────────────────────
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 4

        # ── Respuesta HTTP ───────────────────────────────────────────────────
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="plantilla_clientes.xlsx"'
        wb.save(response)
        return response


# ─────────────────────────────────────────────────────────────────────────────
# Vistas CRUD existentes (sin cambios)
# ─────────────────────────────────────────────────────────────────────────────

class DetalleCliente(LoginRequiredMixin, PermissionRequiredMixin, DetailView):
    model = Clientes
    template_name = 'gestion_clientes/detalle_cliente.html'
    permission_required = 'gestion_clientes.detalle_cliente'
    success_url = reverse_lazy('clientes')


class CrearCliente(LoginRequiredMixin, PermissionRequiredMixin, SuccessMessageMixin, CreateView):
    model = Clientes
    form_class = ClientesForm
    template_name = 'gestion_clientes/registrar_cliente.html'
    permission_required = 'gestion_clientes.registrar_cliente'
    success_url = reverse_lazy('listar_clientes')
    success_message = "El Cliente fue registrado exitosamente"

    def _clean_all_messages(self, request):
        storage = messages.get_messages(request)
        for _ in storage:
            pass
        storage.used = True

    def add_error_messages(self, request, form=None, general_message=None):
        self._clean_all_messages(request)
        if general_message:
            messages.error(request, general_message)
            return
        if form and hasattr(form, 'errors'):
            for field, errors in form.errors.items():
                if errors:
                    messages.error(request, f"Error en {field}: {errors[0]}")
                    return

    def add_success_message(self, request, message):
        self._clean_all_messages(request)
        messages.success(request, message)

    def form_invalid(self, form):
        self.add_error_messages(
            request=self.request,
            form=form,
            general_message="No se pudo registrar el Cliente. Por favor, revise los datos."
        )
        return super().form_invalid(form)

    def form_valid(self, form):
        self._clean_all_messages(self.request)
        response = super().form_valid(form)
        if self.object.pk:
            self.add_success_message(self.request, "Cliente registrado correctamente.")
        return response


class ActualizarCliente(LoginRequiredMixin, PermissionRequiredMixin, SuccessMessageMixin, UpdateView):
    model = Clientes
    form_class = ClientesForm
    template_name = 'gestion_clientes/registrar_cliente.html'
    permission_required = 'gestion_clientes.registrar_cliente'
    success_url = reverse_lazy('listar_clientes')
    success_message = "El Cliente fue actualizado exitosamente"

    def _clean_all_messages(self, request):
        storage = messages.get_messages(request)
        for _ in storage:
            pass
        storage.used = True

    def add_error_messages(self, request, form=None, general_message=None):
        self._clean_all_messages(request)
        if general_message:
            messages.error(request, general_message)
            return
        if form and hasattr(form, 'errors'):
            for field, errors in form.errors.items():
                if errors:
                    messages.error(request, f"Error en {field}: {errors[0]}")
                    return

    def add_success_message(self, request, message):
        self._clean_all_messages(request)
        messages.success(request, message)

    def form_invalid(self, form):
        self.add_error_messages(
            request=self.request,
            form=form,
            general_message="No se pudo registrar el Cliente. Por favor, revise los datos."
        )
        return super().form_invalid(form)

    def form_valid(self, form):
        self._clean_all_messages(self.request)
        response = super().form_valid(form)
        if self.object.pk:
            self.add_success_message(self.request, "Cliente actualizado correctamente.")
        return response


class EliminarCliente(LoginRequiredMixin, PermissionRequiredMixin, DeleteView):
    model = Clientes
    template_name = 'gestion_clientes/eliminar_cliente.html'
    permission_required = 'gestion_clientes.eliminar_cliente'
    success_url = reverse_lazy('listar_clientes')

    def form_valid(self, form):
        try:
            nombre = self.object.nombre
            result = super().form_valid(form)
            messages.success(self.request, f"El Cliente {nombre} fue eliminado exitosamente")
            return result
        except ProtectedError:
            messages.error(
                self.request,
                f"No se pudo eliminar el Cliente {self.object.nombre}. Tiene facturas asociadas."
            )
            return HttpResponseRedirect(self.success_url)
        except Exception as e:
            messages.error(self.request, f"Error inesperado al eliminar el Cliente: {str(e)}")
            return HttpResponseRedirect(self.success_url)